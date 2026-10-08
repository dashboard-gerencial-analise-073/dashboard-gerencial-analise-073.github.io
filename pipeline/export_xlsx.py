"""Exporta a base consolidada (cadastro + CVM + indicadores + qualidade) para Excel."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HDR_FILL = PatternFill("solid", fgColor="0E2A33")
HDR_FONT = Font(bold=True, color="FFFFFF")
PCT = "0.00%"
MONEY = '#,##0'


def _sheet(wb, title, header, rows, formats=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "B2"
    for i, h in enumerate(header, start=1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = min(max(12, len(str(h)) + 2), 48)
        fmt = (formats or {}).get(h)
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                for c in row:
                    c.number_format = fmt
    return ws


def export(ds, path, carteiras=None):
    wb = Workbook()
    wb.remove(wb.active)
    wins = ds["meta"]["janelas"]
    rw = ds["meta"]["janelas_risco"]
    bnames = {b["id"]: b["nome"] for b in ds["benchmarks"]}

    header = ["id", "Fundo", "CNPJ", "Categoria", "Família", "Estrutura", "Classificação ANBIMA (CVM)", "Denominação CVM",
              "Situação CVM", "Administrador", "Gestor", "Início", "Benchmark", "Tx adm (% a.a.)", "Tx perf (%)",
              "Cotização aplic. (D+)", "Cotização resg. (D+)", "Liquidação resg. (D+)", "Prazo", "Qualificado",
              "Previdência", "Isento IR PF", "Data-base", "PL (R$)", "Cotistas", "Captação líq. 12M (R$)", "PL médio 12M (R$)"]
    fmts = {"PL (R$)": MONEY, "Captação líq. 12M (R$)": MONEY, "PL médio 12M (R$)": MONEY}
    for w in wins:
        header += [f"Ret {w['label']}", f"Bench {w['label']}", f"Excesso {w['label']}"]
        fmts.update({f"Ret {w['label']}": PCT, f"Bench {w['label']}": PCT, f"Excesso {w['label']}": PCT})
        if w.get("n", 0) >= 12:
            header += [f"Ret {w['label']} a.a."]
            fmts[f"Ret {w['label']} a.a."] = PCT
    for r in rw:
        header += [f"Vol {r}", f"Sharpe {r}", f"Máx DD {r}", f"Beta {r}", f"TE {r}"]
        fmts.update({f"Vol {r}": PCT, f"Máx DD {r}": PCT, f"TE {r}": PCT, f"Sharpe {r}": "0.00", f"Beta {r}": "0.00"})
    rows = []
    for f in ds["fundos"]:
        cvm, fl = f.get("cvm", {}), f.get("flow", {})
        row = [f["id"], f["nome"], f["cnpj"], f["categoria"], f.get("familia"), f.get("estrutura"),
               cvm.get("classificacao_anbima"), cvm.get("denominacao"), cvm.get("situacao"), cvm.get("administrador"),
               cvm.get("gestor"), f.get("inicio"), bnames.get(f["benchmark_id"]), f.get("taxa_adm"), f.get("taxa_perf"),
               f.get("cot_apl"), f.get("cot_resg"), f.get("liq_resg"), f.get("prazo_tipo"),
               "Sim" if f.get("qualificado") else "Não", "Sim" if f.get("previdencia") else "Não",
               "Sim" if f.get("isento_ir") else "Não", fl.get("data"), fl.get("pl"), fl.get("cotistas"),
               fl.get("captacao_liquida_12m"), fl.get("pl_medio_12m")]
        for w in wins:
            r = f["ret"].get(w["id"], {})
            row += [r.get("f"), r.get("b"), r.get("x")]
            if w.get("n", 0) >= 12:
                row += [r.get("fa")]
        for k in rw:
            r = f["risk"].get(k, {})
            row += [r.get("vol"), r.get("sharpe"), r.get("mdd"), r.get("beta"), r.get("te")]
        rows.append(row)
    _sheet(wb, "Fundos", header, rows, fmts)

    months = sorted({m for f in ds["fundos"] for m in f.get("mensal", {})})
    _sheet(wb, "Rentab_Mensal", ["Fundo"] + months,
           [[f["nome"]] + [f.get("mensal", {}).get(m, [None])[0] for m in months] for f in ds["fundos"]],
           {m: PCT for m in months})

    eixo = ds["eixo"]
    cot_rows = []
    for i, d in enumerate(eixo):
        row = [d]
        for f in ds["fundos"]:
            s = f.get("serie", {})
            j = i - s.get("i0", 10 ** 9)
            row.append(s["q"][j] if 0 <= j < len(s.get("q", [])) else None)
        cot_rows.append(row)
    _sheet(wb, "Cotas_Diarias", ["Data"] + [f["nome"] for f in ds["fundos"]], cot_rows)

    if carteiras:
        rows = [r for r in carteiras if r[10] is None or abs(r[10]) >= 1e-5]  # omite posições < 0,001% do PL
        ws = _sheet(wb, "Carteiras", ["Fundo", "Data carteira", "Grupo", "Tipo de aplicação (CVM)", "Tipo de ativo", "Ativo",
                                      "Código", "Emissor", "Vencimento", "Valor (R$)", "% PL"], rows,
                    {"Valor (R$)": MONEY, "% PL": PCT})
        ws.column_dimensions["F"].width = 60
        ws.auto_filter.ref = ws.dimensions

    q = [[i["nivel"], i.get("fundo_id"), i["codigo"], i["mensagem"]] for i in ds["qualidade"]]
    ws = _sheet(wb, "Qualidade", ["Nível", "Fundo", "Código", "Mensagem"], q)
    ws.column_dimensions["D"].width = 140

    meta = ds["meta"]
    info = [["Data-base", meta["data_base"]], ["Gerado em", meta["gerado_em"]],
            ["Arquivos CVM", meta["cvm_arquivos"]], ["Última data CDI", meta.get("cdi_ultima_data")],
            ["Metodologia", "Ver docs/METODOLOGIA.md"]]
    info += [[s["nome"], s["url"]] for s in meta["fontes"]]
    ws = _sheet(wb, "Fontes", ["Item", "Valor"], info)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 80
    wb.move_sheet("Fontes", offset=-(len(wb.sheetnames) - 1))
    wb.save(path)
    return path
