"""Leitura da planilha-base (data/input/cadastro_fundos.xlsx)."""
from openpyxl import load_workbook

import config

REQUIRED = ["id", "gestora_id", "nome", "cnpj", "categoria", "benchmark_id"]
NUMERIC = ["cotizacao_aplicacao_d", "cotizacao_resgate_d", "liquidacao_resgate_d",
           "taxa_adm_pct", "taxa_perf_pct", "risco_lamina"]
FLAGS = ["ativo", "publico_qualificado", "versao_previdencia", "isento_ir_pf"]


def _sheet(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        rec = {}
        for h, v in zip(header, r):
            if not h:
                continue
            rec[h] = v.strip() if isinstance(v, str) else v
            if rec[h] == "":
                rec[h] = None
        out.append(rec)
    return out


def load(path=config.INPUT_XLSX):
    wb = load_workbook(path, data_only=True, read_only=True)
    gestoras = _sheet(wb, "Gestoras")
    benchmarks = _sheet(wb, "Benchmarks")
    fundos = _sheet(wb, "Fundos")
    errors = []
    for i, f in enumerate(fundos, start=2):
        for k in REQUIRED:
            if f.get(k) in (None, ""):
                errors.append(f"Linha {i}: campo obrigatório '{k}' vazio")
        for k in NUMERIC:
            v = f.get(k)
            if v is not None and not isinstance(v, (int, float)):
                try:
                    f[k] = float(str(v).replace(",", "."))
                except ValueError:
                    errors.append(f"Linha {i} ({f.get('nome')}): '{k}' não numérico: {v!r}")
                    f[k] = None
        for k in FLAGS:
            f[k] = bool(int(f.get(k) or 0)) if k != "ativo" else f.get(k) in (None, 1, "1", True)
    fundos = [f for f in fundos if f["ativo"]]
    return gestoras, benchmarks, fundos, errors
